import openpyxl as xl
from openpyxl.chart import Reference, BarChart

wb = xl.load_workbook('transaction.xlsx')

sheet = wb['Sheet1']

for row in range(2, sheet.max_row +1):
    cell = sheet.cell(row,2)
    corrected_price = cell.value * .9
    corrected_price_cell = sheet.cell(row,3)
    corrected_price_cell.value = corrected_price

values = Reference(
    sheet,
    min_row=2,
    max_row=sheet.max_row,
    min_col=3,
    max_col=3,
)

chart = BarChart()

chart.add_data(values)
sheet.add_chart(chart, 'd1')

wb.save('transaction2.xlsx')
